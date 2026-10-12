from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "APRESENTACAO_COMERCIAL.md"
OUTPUT = ROOT / "docs" / "Proposta Comercial - SPCasting.docx"

INK = "241A16"
ACCENT = "A75F3C"
SOFT = "F4E7DE"
MUTED = "756963"
WHITE = "FFFFFF"


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, value=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    properties.append(margins)


def add_bottom_border(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:color"), ACCENT)
    borders.append(bottom)
    properties.append(borders)


def add_inline(paragraph, text):
    for part in re.split(r"(\*\*.*?\*\*|\[.*?\])", text):
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        value = part[2:-2] if is_bold else part
        run = paragraph.add_run(value)
        run.bold = is_bold
        if value.startswith("[") and value.endswith("]"):
            run.italic = True
            run.font.color.rgb = RGBColor(167, 95, 60)


def add_scope_table(document):
    rows = [
        ("Página institucional", "Apresentação da agência, valores, imagens e acesso à plataforma"),
        ("Cadastro e login", "Acesso por e-mail e senha ou conta Google"),
        ("Perfil do freelancer", "Dados profissionais, cidades, foto, Instagram, banco, PIX e composite"),
        ("Oportunidades", "Busca, detalhes, compatibilidade por cidade e candidatura"),
        ("Inscrições", "Acompanhamento e cancelamento de candidaturas"),
        ("Painel administrativo", "Visão geral de vagas, inscrições e profissionais selecionados"),
        ("Gestão de vagas", "Criação, edição, encerramento e controle de quantidade"),
        ("Gestão de candidatos", "Consulta de perfis, seleção, remoção e contato pelo WhatsApp"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Área"
    table.rows[0].cells[1].text = "O que está incluído"
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        shade_cell(cell, ACCENT)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    for index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if index % 2 == 0:
                shade_cell(cell, SOFT)
    return table


def add_investment_table(document):
    rows = [
        ("Desenvolvimento da plataforma", "Escopo funcional apresentado neste documento", "[preencher]", "Pagamento único"),
        ("Domínio", "Endereço próprio da SPCasting", "[preencher]", "Anual"),
        ("Firebase e publicação", "Serviços utilizados para autenticação, dados e hospedagem", "[preencher]", "Conforme consumo/plano"),
        ("Manutenção", "Suporte, atualizações e pequenas alterações após a entrega", "[preencher]", "Opcional"),
    ]
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, ("Item", "Descrição", "Valor", "Condição")):
        cell.text = value
        shade_cell(cell, ACCENT)
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for index, values in enumerate(rows, start=1):
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if index % 2 == 0:
                shade_cell(cell, SOFT)
        if values[2] == "[preencher]":
            cells[2].paragraphs[0].runs[0].italic = True
            cells[2].paragraphs[0].runs[0].font.color.rgb = RGBColor(167, 95, 60)
    return table


def add_heading(document, text):
    paragraph = document.add_paragraph(style="Heading 2")
    paragraph.add_run(text)
    return paragraph


def add_bullets(document, items):
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def build_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(36, 26, 22)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Aptos Display"
    heading1.font.size = Pt(24)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(36, 26, 22)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Aptos Display"
    heading2.font.size = Pt(15)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(167, 95, 60)
    heading2.paragraph_format.space_before = Pt(11)
    heading2.paragraph_format.space_after = Pt(5)

    title = document.add_paragraph(style="Heading 1")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Proposta Comercial SPCasting")
    date = document.add_paragraph("12 de outubro de 2026")
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bottom_border(document.add_paragraph())

    add_heading(document, "O projeto")
    document.add_paragraph(
        "Uma plataforma web exclusiva para a SPCasting organizar sua rede de talentos, publicar vagas temporárias, receber candidaturas e selecionar profissionais em um único ambiente."
    )

    add_heading(document, "O que será entregue")
    add_scope_table(document)

    add_heading(document, "Também incluído")
    add_bullets(document, [
        "Interface responsiva para celular, tablet e computador.",
        "Cadastro e login com e-mail e senha ou Google.",
        "Dados, fotos e composites privados com acesso controlado.",
        "Integração com Firebase, API de localidades do IBGE e WhatsApp.",
        "Configuração técnica para publicação no endereço definido para o projeto.",
        "Homologação dos principais fluxos antes da entrega.",
    ])

    add_heading(document, "Benefícios para a SPCasting")
    add_bullets(document, [
        "Base de profissionais centralizada e atualizada.",
        "Menos dependência de conversas e planilhas para organizar o casting.",
        "Publicação de oportunidades e formação de equipes com mais rapidez.",
        "Triagem por cidade e disponibilidade das informações profissionais.",
        "Contato direto com candidatos pelo WhatsApp.",
    ])

    add_heading(document, "Investimento")
    add_investment_table(document)
    paragraph = document.add_paragraph()
    add_inline(paragraph, "**Condição de pagamento:** [preencher]")
    paragraph = document.add_paragraph()
    add_inline(paragraph, "**Prazo de entrega:** [preencher]")
    paragraph = document.add_paragraph()
    add_inline(paragraph, "**Validade da proposta:** [preencher]")
    document.add_paragraph("Novas funcionalidades, integrações ou alterações fora deste escopo serão avaliadas e orçadas separadamente.")

    add_heading(document, "Informações necessárias para a entrega")
    add_bullets(document, [
        "Domínio e ambiente de publicação.",
        "Conta e configurações do Firebase.",
        "Definição da conta administrativa.",
        "Validação dos textos, imagens e dados institucionais.",
        "Confirmação das regras operacionais e comerciais das vagas.",
        "Política de privacidade e termos aplicáveis à operação.",
    ])

    add_heading(document, "Próximos passos")
    for item in (
        "Aprovação da proposta e das condições comerciais.",
        "Confirmação dos dados e acessos necessários.",
        "Homologação dos fluxos de freelancer e administrador.",
        "Ajustes finais acordados entre as partes.",
        "Publicação e entrega oficial da plataforma.",
    ):
        document.add_paragraph(item, style="List Number")

    add_bottom_border(document.add_paragraph())
    document.add_paragraph("Responsável pela proposta: [preencher]")
    document.add_paragraph("Contato: [preencher]")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("SPCasting • Proposta comercial")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(117, 105, 99)

    document.core_properties.title = "Proposta Comercial - SPCasting"
    document.core_properties.subject = "Plataforma web de casting e gestão de oportunidades"
    document.core_properties.author = "SPCasting"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
